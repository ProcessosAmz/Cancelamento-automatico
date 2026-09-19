#!/usr/bin/env python3
import json
import unicodedata
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"


def norm(s):
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower().strip()


alvo = json.load(open("saidas/analise_servicos_alvo.json", encoding="utf-8"))
os_map = json.load(open("saidas/analise_os_por_servico.json", encoding="utf-8"))

rows = []
for a in alvo:
    idcs = a["id_cliente_servico"]
    lista = os_map.get(idcs, [])
    qtd_os = len(lista)
    tem_aberta = any(norm(o["status"]) != "finalizado" for o in lista)
    motivos = sorted({o["status_fechamento"] for o in lista if o.get("status_fechamento")})
    tem_motivo_terminal = any(
        any(t in norm(m) for t in ("cancelamento", "inviabilidade")) for m in motivos
    )
    if qtd_os == 0:
        categoria = "Sem nenhuma OS"
    elif tem_aberta:
        categoria = "Tem OS em aberto"
    else:
        categoria = "Tem OS, mas nenhuma em aberto"
    bate_criterio = (categoria != "Tem OS em aberto") or tem_motivo_terminal
    rows.append(
        [
            a["codigo_cliente"],
            a["nome_razaosocial"],
            a["status"],
            idcs,
            qtd_os,
            categoria,
            ", ".join(motivos) if motivos else "",
            "Sim" if tem_motivo_terminal else "Não",
            "Sim" if bate_criterio else "Não",
        ]
    )

rows.sort(key=lambda r: (r[8] != "Sim", r[5], r[0]))

wb = Workbook()

# ---- Sheet Detalhe ----
ws = wb.active
ws.title = "Detalhe"
headers = [
    "Código cliente",
    "Nome / Razão social",
    "Status do serviço",
    "ID cliente_servico",
    "Qtd. OS vinculadas",
    "Categoria",
    "Motivo(s) de fechamento da OS",
    "Tem OS cancelamento/inviabilidade",
    "Bate no critério pedido",
]
for c, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = Font(name=FONT, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="4472C4")
    cell.alignment = Alignment(horizontal="center", wrap_text=True)

for r, row in enumerate(rows, start=2):
    for c, val in enumerate(row, start=1):
        cell = ws.cell(row=r, column=c, value=val)
        cell.font = Font(name=FONT)

widths = [14, 34, 30, 16, 14, 26, 30, 16, 16]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows)+1}"

last_row = len(rows) + 1

# ---- Sheet Resumo ----
ws2 = wb.create_sheet("Resumo", 0)
ws2.column_dimensions["A"].width = 55
ws2.column_dimensions["B"].width = 14

titulo = ws2.cell(row=1, column=1, value="Serviços em Aguardando Assinatura de Contrato / Aguardando Agendamento")
titulo.font = Font(name=FONT, bold=True, size=13)
ws2.merge_cells("A1:B1")

sub = ws2.cell(row=2, column=1, value="sem OS em aberto, ou com OS de cancelamento/inviabilidade")
sub.font = Font(name=FONT, italic=True, size=10)
ws2.merge_cells("A2:B2")

linhas_resumo = [
    ("Total de serviços nesses status (alvo)", f"=COUNTA(Detalhe!A2:A{last_row})"),
    ("  - Aguardando Assinatura de Contrato", f'=COUNTIF(Detalhe!C2:C{last_row},"Aguardando Assinatura de Contrato")'),
    ("  - Aguardando Agendamento", f'=COUNTIF(Detalhe!C2:C{last_row},"Aguardando Agendamento")'),
    ("", ""),
    ("Sem nenhuma OS cadastrada", f'=COUNTIF(Detalhe!F2:F{last_row},"Sem nenhuma OS")'),
    ("Tem OS, mas nenhuma em aberto", f'=COUNTIF(Detalhe!F2:F{last_row},"Tem OS, mas nenhuma em aberto")'),
    ("Tem OS em aberto (excluído, salvo se cancelamento/inviabilidade)", f'=COUNTIF(Detalhe!F2:F{last_row},"Tem OS em aberto")'),
    ("Com OS marcada como cancelamento/inviabilidade", f'=COUNTIF(Detalhe!H2:H{last_row},"Sim")'),
    ("", ""),
    ("TOTAL que bate no critério pedido", f'=COUNTIF(Detalhe!I2:I{last_row},"Sim")'),
    ("% do total alvo", f"=B13/B4"),
]

r = 4
for label, formula in linhas_resumo:
    la = ws2.cell(row=r, column=1, value=label)
    la.font = Font(name=FONT, bold=label.startswith("TOTAL"))
    if formula:
        vb = ws2.cell(row=r, column=2, value=formula)
        vb.font = Font(name=FONT, bold=label.startswith("TOTAL"))
        if "%" in label:
            vb.number_format = "0.0%"
    r += 1

nota_r = r + 1
nota = ws2.cell(
    row=nota_r,
    column=1,
    value=(
        "Nota importante: a API do HubSoft não registra 'cancelamento' e 'inviabilidade' como "
        "valores distintos no campo de motivo de fechamento da OS - só existem 'concluido' e "
        "'sem_conclusao'. Por isso a contagem específica de cancelamento/inviabilidade deu 0; "
        "esses casos provavelmente estão dentro de 'sem_conclusao' (fechada sem concluir o "
        "serviço), que já está incluído em 'Tem OS, mas nenhuma em aberto'. Fonte: HubSoft API "
        "GraphQL (/graphql/v1), consulta feita em 18/09/2026."
    ),
)
nota.font = Font(name=FONT, italic=True, size=9, color="808080")
nota.alignment = Alignment(wrap_text=True)
ws2.merge_cells(start_row=nota_r, start_column=1, end_row=nota_r, end_column=2)
ws2.row_dimensions[nota_r].height = 60

wb.save("planilha_servicos_pendentes.xlsx")
print("salvo: planilha_servicos_pendentes.xlsx", len(rows), "linhas")